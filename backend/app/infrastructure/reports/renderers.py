"""Eksport renderer'lari (BE-602).

- CSV: faqat tranzaksiyalar.
- XLSX: 2 varaq — Summary (Income/Expenses/Net + kategoriyalar), Transactions.
- PDF: sarlavha, davr, Income/Expenses/Net, kategoriya jadvali, tranzaksiyalar ro'yxati.
CSV/XLSX'da `=`, `+`, `-`, `@` bilan boshlanuvchi matn hujayralari oldiga `'` qo'yiladi.
"""

import io
import logging
from pathlib import Path

from app.application.common.interfaces import RenderedReport, ReportRenderer
from app.application.exports.csv_safe import render_csv, sanitize_cell
from app.application.exports.report import ExportReport
from app.core.i18n import format_amount
from app.domain.exports.entities import CONTENT_TYPES, ExportFormat

logger = logging.getLogger("finora.exports")


def _columns(report: ExportReport) -> list[str]:
    lb = report.labels
    return [lb.col_date, lb.col_type, lb.col_category, lb.col_title, lb.col_amount, lb.col_note]


class CsvRenderer:
    def render(self, report: ExportReport) -> RenderedReport:
        # CSV'da summa musbat (yo'nalish "Type" ustunida) — manfiy son "-" bilan boshlanib
        # formula-himoyasi ostida "'-5000" bo'lib qolmasin
        data = render_csv(_columns(report), (
            [r.date, r.type, r.category, r.title, abs(r.amount), r.note] for r in report.rows))
        return RenderedReport(data, CONTENT_TYPES[ExportFormat.CSV])


class XlsxRenderer:
    def render(self, report: ExportReport) -> RenderedReport:
        from openpyxl import Workbook
        from openpyxl.styles import Font

        lb = report.labels
        wb = Workbook()
        summary = wb.active
        assert summary is not None
        summary.title = lb.summary[:31]
        bold = Font(bold=True)
        summary.append([sanitize_cell(lb.title)])
        summary["A1"].font = Font(bold=True, size=14)
        summary.append([sanitize_cell(lb.period), sanitize_cell(report.range_label)])
        summary.append([])
        for label, value in ((lb.income, report.income), (lb.expenses, report.expenses),
                             (lb.net, report.net)):
            summary.append([sanitize_cell(label), value])
        summary.append([])
        summary.append([sanitize_cell(lb.by_category)])
        summary.cell(row=summary.max_row, column=1).font = bold
        summary.append([sanitize_cell(lb.col_category), sanitize_cell(lb.col_amount),
                        sanitize_cell(lb.col_share)])
        for c in report.categories:
            summary.append([sanitize_cell(c.name), c.amount, c.pct / 100])
            summary.cell(row=summary.max_row, column=3).number_format = "0%"
        summary.column_dimensions["A"].width = 28
        summary.column_dimensions["B"].width = 18

        txs = wb.create_sheet(lb.transactions[:31])
        txs.append([sanitize_cell(c) for c in _columns(report)])
        for cell in txs[1]:
            cell.font = bold
        for r in report.rows:
            # Summa — son (formula emas); matnlar tozalangan
            txs.append([r.date, sanitize_cell(r.type), sanitize_cell(r.category),
                        sanitize_cell(r.title), r.amount, sanitize_cell(r.note)])
        for col, width in zip("ABCDEF", (12, 12, 20, 28, 16, 32), strict=True):
            txs.column_dimensions[col].width = width
        buf = io.BytesIO()
        wb.save(buf)
        return RenderedReport(buf.getvalue(), CONTENT_TYPES[ExportFormat.XLSX])


_FALLBACK = str.maketrans({"ʻ": "'", "ʼ": "'", "‘": "'", "’": "'", "–": "-", "—": "-",
                           "…": "...", "₽": "RUB", "₸": "KZT", "₺": "TRY"})


class PdfRenderer:
    """fpdf2. Unicode shrift (FINORA_PDF_FONT_PATH, Dockerda DejaVuSans) bo'lmasa —
    Helvetica va lotin-1 ga moslashtirilgan matn (kirill harflari '?' bo'lib qoladi)."""

    def __init__(self, font_path: str | None) -> None:
        self._font = font_path if font_path and Path(font_path).is_file() else None
        if font_path and self._font is None:
            logger.warning("pdf_font_missing")

    def _text(self, value: str) -> str:
        if self._font:
            return value
        return value.translate(_FALLBACK).encode("latin-1", "replace").decode("latin-1")

    def render(self, report: ExportReport) -> RenderedReport:
        from fpdf import FPDF

        pdf = FPDF(format="A4")
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()
        family = "Helvetica"
        if self._font:
            pdf.add_font("Body", "", self._font)
            pdf.add_font("Body", "B", self._font)
            family = "Body"
        lb, tx = report.labels, self._text

        pdf.set_font(family, "B", 18)
        pdf.cell(0, 10, tx(lb.title), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font(family, "", 11)
        pdf.cell(0, 7, tx(f"{lb.period}: {report.range_label}"), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)
        for label, value in ((lb.income, report.income), (lb.expenses, report.expenses),
                             (lb.net, report.net)):
            pdf.set_font(family, "B", 11)
            pdf.cell(45, 7, tx(label))
            pdf.set_font(family, "", 11)
            pdf.cell(0, 7, f"{format_amount(value)} UZS", new_x="LMARGIN", new_y="NEXT")

        if report.categories:
            pdf.ln(4)
            pdf.set_font(family, "B", 13)
            pdf.cell(0, 8, tx(lb.by_category), new_x="LMARGIN", new_y="NEXT")
            pdf.set_font(family, "", 10)
            for c in report.categories:
                pdf.cell(90, 6, tx(c.name[:48]))
                pdf.cell(50, 6, format_amount(c.amount), align="R")
                pdf.cell(0, 6, f"{c.pct}%", align="R", new_x="LMARGIN", new_y="NEXT")

        pdf.ln(4)
        pdf.set_font(family, "B", 13)
        pdf.cell(0, 8, tx(lb.transactions), new_x="LMARGIN", new_y="NEXT")
        widths = (24, 24, 40, 56, 36)
        pdf.set_font(family, "B", 9)
        for w, h in zip(widths, (lb.col_date, lb.col_type, lb.col_category, lb.col_title,
                                 lb.col_amount), strict=True):
            pdf.cell(w, 6, tx(h[:24]), border="B")
        pdf.ln()
        pdf.set_font(family, "", 9)
        for r in report.rows:
            cells = (r.date, r.type[:14], r.category[:24], r.title[:34], format_amount(r.amount))
            for i, (w, text) in enumerate(zip(widths, cells, strict=True)):
                pdf.cell(w, 5.5, tx(text), align="R" if i == 4 else "L")
            pdf.ln()
        return RenderedReport(bytes(pdf.output()), CONTENT_TYPES[ExportFormat.PDF])


def build_renderers(font_path: str | None) -> dict[ExportFormat, ReportRenderer]:
    return {ExportFormat.CSV: CsvRenderer(), ExportFormat.XLSX: XlsxRenderer(),
            ExportFormat.PDF: PdfRenderer(font_path)}
