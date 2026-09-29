"""Goals (BE-1101..1105), insights (BE-901/902), eksport (BE-601/602), yordam (BE-1701/1702)."""

from datetime import UTC, datetime

from tests.conftest import expense, idem, income, login

PHONE = "+998908880001"


class FixedClock:
    def __init__(self, at: datetime) -> None:
        self.at = at

    def now(self) -> datetime:
        return self.at


async def _goal(client, s, **body):
    r = await client.post("/v1/goals", headers={**s.headers, **idem()},
                          json={"name": "Mashina", "target": 1_200_000, **body})
    assert r.status_code == 201, r.text
    return r.json()


async def test_goal_validation_codes(client, container):
    s = await login(client, container, PHONE)
    r = await client.post("/v1/goals", headers={**s.headers, **idem()},
                          json={"name": "  ", "target": 500_000})
    assert r.status_code == 422 and r.json()["code"] == "goal_name_required"
    r = await client.post("/v1/goals", headers={**s.headers, **idem()},
                          json={"name": "Phone", "target": 99_999})
    assert r.status_code == 422 and r.json()["code"] == "goal_target_min"


async def test_goal_plan_history_and_delete_returns_money(client, container):
    container.clock = FixedClock(datetime(2026, 9, 28, 7, 0, tzinfo=UTC))
    s = await login(client, container, PHONE)
    await client.post("/v1/onboarding/balance", headers=s.headers,
                      json={"amount": 2_000_000, "location": "cash"})
    goal = await _goal(client, s, deadline="2027-03-28", auto_save_monthly=150_000)
    assert goal["plan"] == {"remaining": 1_200_000, "monthly_needed": 200_000,
                            "eta_months": 8}
    for amount in (300_000, 400_000):
        await client.post(f"/v1/goals/{goal['id']}/deposits", headers={**s.headers, **idem()},
                          json={"amount": amount})
    await client.post(f"/v1/goals/{goal['id']}/withdrawals", headers={**s.headers, **idem()},
                      json={"amount": 100_000})
    g = (await client.get(f"/v1/goals/{goal['id']}", headers=s.headers)).json()
    assert (g["saved"], g["pct"]) == (600_000, 50)
    hist = (await client.get(f"/v1/goals/{goal['id']}/history?limit=2",
                             headers=s.headers)).json()
    assert [e["kind"] for e in hist["items"]] == ["withdraw", "deposit"] and hist["next_cursor"]
    assert (await client.get("/v1/home", headers=s.headers)).json()["balance"]["total"] == \
        1_400_000

    r = await client.delete(f"/v1/goals/{goal['id']}", headers=s.headers)
    assert r.status_code == 200 and r.json() == {"returned_amount": 600_000}
    assert (await client.get("/v1/home", headers=s.headers)).json()["balance"]["total"] == \
        2_000_000
    assert (await client.get(f"/v1/goals/{goal['id']}", headers=s.headers)).status_code == 404

    plan = (await client.get("/v1/goals/plan?target=1000000&auto_save=300000",
                             headers=s.headers)).json()
    assert plan == {"remaining": 1_000_000, "monthly_needed": None, "eta_months": 4}


async def test_goal_milestones_notified_once(client, container):
    s = await login(client, container, PHONE)
    goal = await _goal(client, s, target=100_000)
    for amount in (30_000, 30_000, 40_000):
        await client.post(f"/v1/goals/{goal['id']}/deposits", headers={**s.headers, **idem()},
                          json={"amount": amount})
    items = (await client.get("/v1/notifications", headers=s.headers)).json()["items"]
    titles = sorted(i["title"] for i in items if i["type"] == "goal_milestone")
    assert titles == ["Mashina: 100% ga yetdingiz", "Mashina: 25% ga yetdingiz",
                      "Mashina: 50% ga yetdingiz", "Mashina: 75% ga yetdingiz"]


async def test_autosave_job_is_idempotent_and_skips_without_money(client, container):
    from app.jobs import run

    container.clock = FixedClock(datetime(2026, 9, 5, 7, 0, tzinfo=UTC))
    s = await login(client, container, PHONE)
    await client.post("/v1/onboarding/balance", headers=s.headers,
                      json={"amount": 250_000, "location": "cash"})
    goal = await _goal(client, s, auto_save_monthly=200_000, auto_save_day=3)
    assert (await run(container, "autosave"))["autosaved"] == 1
    assert (await run(container, "autosave"))["autosaved"] == 0  # shu oy allaqachon
    g = (await client.get(f"/v1/goals/{goal['id']}", headers=s.headers)).json()
    assert g["saved"] == 200_000
    container.clock = FixedClock(datetime(2026, 10, 5, 7, 0, tzinfo=UTC))
    result = await run(container, "autosave")
    assert (result["autosaved"], result["skipped"]) == (0, 1)  # hisobda 50 000 qoldi
    items = (await client.get("/v1/notifications", headers=s.headers)).json()["items"]
    assert any("mablag' yetmadi" in i["body"] for i in items)


async def test_insights_detect_act_and_dismiss(client, container):
    container.clock = FixedClock(datetime(2026, 9, 20, 7, 0, tzinfo=UTC))
    s = await login(client, container, PHONE)
    for month in ("06", "07", "08"):
        await expense(client, s, 300_000, "food", occurred_at=f"2026-{month}-10T12:00:00+05:00")
    await expense(client, s, 600_000, "food", occurred_at="2026-09-10T12:00:00+05:00")
    for i in range(8):
        await expense(client, s, 50_000, "groceries",
                      occurred_at=f"2026-09-{i + 1:02d}T12:00:00+05:00")
    for title, amount in (("Netflix", 120_000), ("Spotify", 60_000)):
        await expense(client, s, amount, "subs", title=title,
                      occurred_at="2026-09-03T12:00:00+05:00")
    await income(client, s, 5_000_000, occurred_at="2026-09-01T12:00:00+05:00")
    await _goal(client, s)

    data = (await client.get("/v1/insights", headers=s.headers)).json()
    by_kind = {i["kind"]: i for i in data["items"]}
    assert set(by_kind) == {"overspend", "discount_days", "subscriptions", "autosave"}
    over = by_kind["overspend"]
    assert over["category_id"] == "food" and over["action"] == "set_budget"
    assert over["saving"] == 300_000 and "100%" in over["title"]
    assert by_kind["subscriptions"]["saving"] == 60_000
    assert by_kind["autosave"]["saving"] == 500_000
    assert data["potential_saving"] == sum(i["saving"] for i in data["items"])
    home = (await client.get("/v1/home", headers=s.headers)).json()
    assert home["ai_teaser"]["saving"] == data["potential_saving"]

    r = await client.post(f"/v1/insights/{over['id']}/action", headers=s.headers, json={})
    assert r.json() == {"status": "done", "category_id": "food", "monthly_limit": 300_000,
                        "reminder_id": None, "goal_id": None, "auto_save_monthly": None}
    food = next(c for c in (await client.get("/v1/categories", headers=s.headers)).json()
                if c["id"] == "food")
    assert food["monthly_limit"] == 300_000

    r = await client.post(f"/v1/insights/{by_kind['discount_days']['id']}/action",
                          headers=s.headers)
    assert r.json()["reminder_id"]
    r = await client.post(f"/v1/insights/{by_kind['autosave']['id']}/action",
                          headers=s.headers)
    assert r.json()["auto_save_monthly"] == 500_000
    sub_id = by_kind["subscriptions"]["id"]
    assert (await client.post(f"/v1/insights/{sub_id}/dismiss",
                              headers=s.headers)).status_code == 204
    assert (await client.get("/v1/insights", headers=s.headers)).json()["items"] == []
    shown = (await client.get("/v1/insights?include_dismissed=true",
                              headers=s.headers)).json()["items"]
    assert {i["status"] for i in shown} == {"done", "dismissed"}

    other = await login(client, container, "+998908880002")
    r = await client.post(f"/v1/insights/{sub_id}/dismiss", headers=other.headers)
    assert r.status_code == 404


async def test_export_preview_names_and_formats(client, container):
    container.clock = FixedClock(datetime(2026, 9, 28, 7, 0, tzinfo=UTC))
    s = await login(client, container, PHONE)
    await expense(client, s, 120_000, "food", title="Evos", occurred_at="2026-09-28T10:00:00+05:00")
    await income(client, s, 500_000, occurred_at="2026-09-02T10:00:00+05:00")
    await expense(client, s, 70_000, "transport", occurred_at="2026-01-15T10:00:00+05:00")

    async def preview(period, include="expense,income,transfer", fmt="pdf"):
        r = await client.get(f"/v1/exports/preview?period={period}&include={include}"
                             f"&format={fmt}", headers=s.headers)
        assert r.status_code == 200, r.text
        return r.json()

    m = await preview("monthly")
    assert (m["count"], m["income"], m["expenses"], m["net"]) == (2, 500_000, 120_000, 380_000)
    assert m["file_name"] == "finora_monthly_sep2026.pdf"
    assert (m["from"], m["to"]) == ("2026-09-01", "2026-09-30")
    assert (await preview("weekly", fmt="xlsx"))["file_name"] == "finora_weekly_w40_2026.xlsx"
    assert (await preview("daily", fmt="csv"))["file_name"] == "finora_daily_28sep2026.csv"
    y = await preview("yearly")
    assert y["file_name"] == "finora_yearly_2026.pdf" and y["expenses"] == 190_000
    only_exp = await preview("monthly", include="expense")
    assert (only_exp["income"], only_exp["count"]) == (0, 1)
    r = await client.get("/v1/exports/preview?period=monthly&include=", headers=s.headers)
    assert r.status_code == 422
    r = await client.post("/v1/exports", headers=s.headers,
                          json={"period": "monthly", "include": [], "format": "pdf"})
    assert r.status_code == 422

    for fmt, magic in (("pdf", b"%PDF"), ("xlsx", b"PK")):
        r = await client.post("/v1/exports", headers=s.headers,
                              json={"period": "yearly", "include": ["expense", "income"],
                                    "format": fmt})
        status = (await client.get(f"/v1/exports/{r.json()['id']}", headers=s.headers)).json()
        assert status["status"] == "ready" and status["row_count"] == 3, status
        f = await client.get(status["download_url"])
        assert f.status_code == 200 and f.content.startswith(magic)
        assert status["file_name"] == f"finora_yearly_2026.{fmt}"


async def test_xlsx_has_two_sheets_and_safe_cells(client, container):
    import io

    from openpyxl import load_workbook

    s = await login(client, container, PHONE)
    await expense(client, s, 1000, "food", title="=HYPERLINK(\"x\")")
    r = await client.post("/v1/exports", headers=s.headers,
                          json={"period": "monthly", "include": ["expense"], "format": "xlsx"})
    url = (await client.get(f"/v1/exports/{r.json()['id']}", headers=s.headers)).json()[
        "download_url"]
    wb = load_workbook(io.BytesIO((await client.get(url)).content))
    assert wb.sheetnames == ["Xulosa", "Tranzaksiyalar"]
    row = [c.value for c in wb["Tranzaksiyalar"][2]]
    assert row[3] == "'=HYPERLINK(\"x\")" and row[4] == -1000


async def test_export_failure_status_and_rate_limit(client, container):
    class Broken:
        def render(self, report):
            raise RuntimeError("boom")

    from app.domain.exports.entities import ExportFormat

    container.renderers[ExportFormat.PDF] = Broken()
    container.settings.exports_per_hour = 2
    s = await login(client, container, PHONE)
    body = {"period": "daily", "include": ["expense"], "format": "pdf"}
    r = await client.post("/v1/exports", headers=s.headers, json=body)
    status = (await client.get(f"/v1/exports/{r.json()['id']}", headers=s.headers)).json()
    assert status["status"] == "failed" and status["error_code"] == "render_failed"
    assert status["download_url"] is None
    assert (await client.post("/v1/exports", headers=s.headers, json=body)).status_code == 202
    assert (await client.post("/v1/exports", headers=s.headers, json=body)).status_code == 429


async def test_help_faq_contacts_and_support_session(client, container):
    from pydantic import SecretStr

    from app.domain.help.entities import FaqItem

    s = await login(client, container, PHONE)
    async with container.uow() as uow:
        await uow.faq.replace_language("en", [FaqItem(__import__("uuid").uuid4(), "en",
                                                      "How do I add a card?", "Tap +", 0)])
        await uow.commit()
    r = await client.get("/v1/help/faq?lang=tr", headers=s.headers)
    assert [i["question"] for i in r.json()] == ["How do I add a card?"]  # en'ga qaytadi
    contacts = (await client.get("/v1/help/contacts", headers=s.headers)).json()
    assert contacts["email"] == "help@finora.uz" and contacts["live_chat"] is False
    assert (await client.post("/v1/support/session", headers=s.headers)).status_code == 503
    container.settings.support_chat_secret = SecretStr("chat-secret")
    body = (await client.post("/v1/support/session", headers=s.headers)).json()
    assert body["user_id"] == s.user["id"] and len(body["user_hash"]) == 64


async def test_weekly_report_on_monday_morning(client, container):
    from app.jobs import run

    container.clock = FixedClock(datetime(2026, 9, 27, 7, 0, tzinfo=UTC))
    s = await login(client, container, PHONE)
    await expense(client, s, 100_000, "food", occurred_at="2026-09-16T12:00:00+05:00")
    await expense(client, s, 150_000, "food", occurred_at="2026-09-23T12:00:00+05:00")
    container.clock = FixedClock(datetime(2026, 9, 28, 3, 0, tzinfo=UTC))  # 08:00 Toshkent
    assert (await run(container, "weekly"))["weekly_sent"] == 0
    container.clock = FixedClock(datetime(2026, 9, 28, 4, 30, tzinfo=UTC))  # 09:30
    assert (await run(container, "weekly"))["weekly_sent"] == 1
    assert (await run(container, "weekly"))["weekly_sent"] == 0
    items = (await client.get("/v1/notifications", headers=s.headers)).json()["items"]
    report = next(i for i in items if i["type"] == "weekly_report")
    assert "50%" in report["body"] and report["deep_link"] == "/stats"
